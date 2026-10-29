"""Build test fixtures from spreadsheets.

  python3 tools/fixtures.py demo     spec/SLAM Build Spec.xlsx, tab 'Demo client' -> fixtures/demo.json
  python3 tools/fixtures.py private  fixtures/SLAM Test Client (private).xlsx -> fixtures/eli.local.json

Expected outputs are spreadsheet formulas. They are worked out here, in Python,
independently of the app's JavaScript, so the tests compare two separate answers.
The private file and its output are git-ignored. Never commit either.
"""
import json, math, pathlib, re, sys
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
DEMO_SRC = ROOT / "spec" / "SLAM Build Spec.xlsx"
PRIVATE_SRC = ROOT / "fixtures" / "SLAM Test Client (private).xlsx"

# Spreadsheet label -> fixture key. Both files use these labels.
INPUT_KEYS = {
    "age": "age", "state": "state", "household": "household", "work situation": "workSituation",
    "monthly take-home": "takeHome", "monthly take-home (new job)": "takeHome",
    "monthly spending, all in": "spendingAllIn", "rent (my share)": "rent", "utilities": "utilities",
    "phone (real cost, covered by parents)": "phoneCoveredByParents",
    "cash savings": "cash", "invested (all accounts)": "invested", "car value": "carValue",
    "car loan balance": "carLoanBalance", "car loan apr": "carLoanApr", "car loan payment / mo": "carLoanPayment",
    "student loan balance": "studentLoanBalance", "student loan apr": "studentLoanApr",
    "student loan payment / mo": "studentLoanPayment",
    "credit card balance": "cardBalance", "credit card apr": "cardApr", "credit card payment / mo": "cardPayment",
    "return": "realReturn", "return (after inflation)": "realReturn", "inflation": "inflation",
    "withdrawal rate": "withdrawalRate",
}
EXPECTED_KEYS = [
    ("total net worth", "netWorth"), ("runway", "runwayMonths"), ("monthly surplus", "surplus"),
    ("savings rate", "savingsRate"), ("housing %", "housingPct"),
    ("credit card payoff months", "cardPayoffMonths"), ("credit card total interest", "cardInterest"),
    ("student loan payoff months", "studentLoanPayoffMonths"), ("student loan total interest", "studentLoanInterest"),
    ("car loan payoff months", "carLoanPayoffMonths"), ("real return used", "realReturn"),
    ("rule of 5 target months", "ruleOf5Months"), ("rule of 5 target $", "ruleOf5Target"),
    ("rule of 5 keep-pace", "ruleOf5KeepPace"), ("coast fi number", "coastFi"), ("fi number", "fiNumber"),
    ("months to fi", "monthsToFi"), ("support received", "supportReceived"),
]
MONTH_KEYS = {"cardPayoffMonths", "studentLoanPayoffMonths", "carLoanPayoffMonths", "monthsToFi", "runwayMonths", "ruleOf5Months"}


class Blank:
    """An empty cell. Arithmetic on it stays blank; it equals ''."""
    def _b(self, *_): return BLANK
    __add__ = __radd__ = __sub__ = __rsub__ = __mul__ = __rmul__ = _b
    __truediv__ = __rtruediv__ = __pow__ = __rpow__ = __neg__ = _b
    def __eq__(self, o): return o == "" or isinstance(o, Blank)
    def __hash__(self): return 0
BLANK = Blank()


def nper(rate, pmt, pv, fv=0):
    if any(isinstance(x, Blank) for x in (rate, pmt, pv, fv)): return BLANK
    if rate == 0: return -(pv + fv) / pmt
    return math.log((pmt - fv * rate) / (pmt + pv * rate)) / math.log(1 + rate)


def roundup(x, digits=0):
    if isinstance(x, Blank): return BLANK
    f = 10 ** digits
    return math.ceil(x * f - 1e-9) / f


FUNCS = {"NPER": nper, "ROUNDUP": roundup, "IF": lambda c, a, b: a if c else b, "OR": lambda *a: any(a)}


def evaluate(formula, cells):
    expr = formula[1:]
    expr = re.sub(r'(?<![<>!=])=(?!=)', '==', expr)
    expr = expr.replace('^', '**').replace('<>', '!=')
    expr = re.sub(r'\b([A-Z]{1,2}\d+)\b(?!\()', lambda m: f'C("{m.group(1)}")', expr)
    def C(ref):
        v = cells.get(ref)
        if v is None: return BLANK
        if isinstance(v, str) and v.startswith("="): return evaluate(v, cells)
        return v
    return eval(expr, {"__builtins__": {}}, {"C": C, **FUNCS})


def read(path, tab=None):
    wb = openpyxl.load_workbook(path)
    ws = wb[tab] if tab else wb.worksheets[0]
    cells = {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
    inputs, expected, notes, section = {}, {}, {}, "inputs"
    for row in ws.iter_rows():
        label, value = row[0].value, row[1].value if len(row) > 1 else None
        note = row[2].value if len(row) > 2 else None
        if not isinstance(label, str): continue
        if label.upper().startswith("EXPECTED"): section = "expected"; continue
        low = label.strip().lower()
        if section == "inputs" and low in INPUT_KEYS:
            key = INPUT_KEYS[low]
            inputs[key] = value  # None = not entered (empty is not zero)
            if note: notes[key] = str(note)
            # A 0% loan with no APR row: the note says so.
            if key.endswith("Balance") and note and re.search(r'\b0% APR\b', str(note)):
                inputs.setdefault(key.replace("Balance", "Apr"), 0)
        elif section == "expected" and isinstance(value, str) and value.startswith("="):
            key = next(  # first match wins, so longer labels come first
                (k for words, k in EXPECTED_KEYS if words in low), None)
            if key is None: raise SystemExit(f"Unknown expected output: {label}")
            v = evaluate(value, cells)
            expected[key] = {"label": label, "value": None if isinstance(v, Blank) else v,
                             "tolerance": "1 month" if key in MONTH_KEYS else "$1"}
            if isinstance(v, str): expected[key] = {"label": label, "value": None, "needs": v}
    return inputs, expected, notes


def main(which):
    if which == "demo":
        inputs, expected, _ = read(DEMO_SRC, "Demo client")
        out = ROOT / "fixtures" / "demo.json"
        doc = {"name": "Demo client", "about": "Made up. Safe for the public repo.", "inputs": inputs, "expected": expected}
    elif which == "private":
        if not PRIVATE_SRC.exists():
            raise SystemExit(f"Missing {PRIVATE_SRC.name}. Put it in fixtures/ (it is git-ignored).")
        inputs, expected, notes = read(PRIVATE_SRC)
        out = ROOT / "fixtures" / "eli.local.json"
        doc = {"name": "Test client", "about": "Private. Never commit.", "inputs": inputs, "notes": notes, "expected": expected}
    else:
        raise SystemExit(__doc__)
    out.write_text(json.dumps(doc, indent=2) + "\n")
    print(f"Wrote {out.relative_to(ROOT)}: {len(doc['inputs'])} inputs, {len(doc['expected'])} expected outputs")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "")
